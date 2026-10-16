"""Uploaded-file abstraction: kind sniffing, encoding detection and tabular reading."""

import codecs
import csv
import datetime as dt
import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal
from functools import cached_property
from html.parser import HTMLParser
from typing import Literal

import chardet
import xlrd
from openpyxl import load_workbook

from app.parsers.text import looks_like_date, looks_numeric

FileKind = Literal["csv", "xlsx", "xls", "html", "xml", "pdf", "unknown"]
TableKind = Literal["csv", "xlsx", "xls", "html"]

DELIMITER_CANDIDATES = (",", ";", "\t", "|")


class UnsupportedFileError(ValueError):
    pass


@dataclass
class Table:
    rows: list[list[str]]
    kind: TableKind
    delimiter: str | None = None
    encoding: str | None = None
    sheet_name: str | None = None
    # Spreadsheet number cells arrive already canonical ("1234.5"), so per-cell inference is safer
    # than applying a file-wide decimal separator.
    typed_cells: bool = False
    _header_cache: dict[int | None, int] = field(default_factory=dict, repr=False)

    def detect_header_row(self) -> int:
        """Index of the header row, skipping title/preamble lines some banks put above the table."""
        if None in self._header_cache:
            return self._header_cache[None]
        widths = Counter(_filled(row) for row in self.rows[:300] if _filled(row) >= 2)
        if not widths:
            self._header_cache[None] = 0
            return 0
        target = widths.most_common(1)[0][0]
        result = 0
        for index, row in enumerate(self.rows[:40]):
            cells = [cell for cell in row if cell.strip()]
            if len(cells) < max(2, int(target * 0.6)):
                continue
            textual = sum(1 for cell in cells if not looks_numeric(cell) and not looks_like_date(cell))
            if textual / len(cells) >= 0.7:
                result = index
                break
        self._header_cache[None] = result
        return result

    def headers(self, header_row: int | None = None) -> list[str]:
        index = self.detect_header_row() if header_row is None else header_row
        if index >= len(self.rows):
            return []
        return unique_headers(self.rows[index])

    def data_rows(self, header_row: int | None = None) -> list[tuple[int, list[str]]]:
        """(1-based source line number, cells) for every non-empty row below the header."""
        index = self.detect_header_row() if header_row is None else header_row
        return [(number + 1, row) for number, row in enumerate(self.rows) if number > index and _filled(row) > 0]


def _filled(row: list[str]) -> int:
    return sum(1 for cell in row if cell.strip())


def unique_headers(raw: list[str]) -> list[str]:
    seen: Counter[str] = Counter()
    headers: list[str] = []
    for position, cell in enumerate(raw):
        name = " ".join(cell.split()) or f"Column {position + 1}"
        seen[name] += 1
        headers.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return headers


class SourceFile:
    """An uploaded statement. Everything is derived lazily from the raw bytes."""

    def __init__(self, filename: str, content: bytes) -> None:
        self.filename = filename or "upload"
        self.content = content

    @cached_property
    def sha256(self) -> str:
        return hashlib.sha256(self.content).hexdigest()

    @cached_property
    def extension(self) -> str:
        return self.filename.rsplit(".", 1)[-1].lower() if "." in self.filename else ""

    @cached_property
    def kind(self) -> FileKind:
        head = self.content[:512]
        if head.startswith(b"PK\x03\x04"):
            return "xlsx"
        if head.startswith(b"\xd0\xcf\x11\xe0"):
            return "xls"
        if head.startswith(b"%PDF"):
            return "pdf"
        try:
            text_head = self.text[:2048].lstrip().lower()
        except UnsupportedFileError:
            return "unknown"
        if text_head.startswith(("<html", "<!doctype html", "<table")) or (
            self.extension in {"xls", "xlsx"} and "<table" in text_head
        ):
            return "html"
        if text_head.startswith("<?xml") or text_head.startswith("<flexqueryresponse"):
            return "xml"
        return "csv"

    @cached_property
    def _decoded(self) -> tuple[str, str]:
        return decode_bytes(self.content)

    @property
    def text(self) -> str:
        return self._decoded[0]

    @property
    def encoding(self) -> str:
        return self._decoded[1]

    @cached_property
    def table(self) -> Table:
        kind = self.kind
        if kind == "csv":
            delimiter = sniff_delimiter(self.text)
            rows = list(csv.reader(io.StringIO(self.text), delimiter=delimiter))
            return Table(rows=rows, kind="csv", delimiter=delimiter, encoding=self.encoding)
        if kind == "xlsx":
            return read_xlsx(self.content)
        if kind == "xls":
            return read_xls(self.content)
        if kind == "html":
            return read_html_table(self.text, self.encoding)
        raise UnsupportedFileError(f"'{self.filename}' is not a tabular file ({kind}).")

    def table_with(self, *, delimiter: str | None = None, encoding: str | None = None) -> Table:
        """Re-read a CSV with user-overridden delimiter/encoding (no-op for spreadsheets)."""
        if self.kind != "csv" or (delimiter is None and encoding is None):
            return self.table
        text = self.content.decode(encoding, errors="replace") if encoding else self.text
        text = text.lstrip("﻿")
        chosen = delimiter or sniff_delimiter(text)
        rows = list(csv.reader(io.StringIO(text), delimiter=chosen))
        return Table(rows=rows, kind="csv", delimiter=chosen, encoding=encoding or self.encoding)


def decode_bytes(content: bytes) -> tuple[str, str]:
    """Decode text, preferring UTF-8 and handling Central European legacy encodings."""
    if content.startswith(codecs.BOM_UTF8):
        return content[len(codecs.BOM_UTF8) :].decode("utf-8", errors="replace"), "utf-8-sig"
    if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return content.decode("utf-16", errors="replace").lstrip("﻿"), "utf-16"
    if b"\x00" in content[:4096]:
        # Binary formats are sniffed by magic bytes before decoding; anything else with NULs is not text.
        if content[:4096].count(b"\x00") > len(content[:4096]) // 4:
            for candidate in ("utf-16-le", "utf-16-be"):
                try:
                    return content.decode(candidate), candidate
                except UnicodeDecodeError:
                    continue
        raise UnsupportedFileError("File looks binary, not text.")
    try:
        return content.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    guess = chardet.detect(content[:200_000]).get("encoding") or "cp1250"
    candidates = [guess.lower(), "cp1250", "cp1252", "latin-1"]
    for candidate in dict.fromkeys(candidates):
        try:
            text = content.decode(candidate)
        except (UnicodeDecodeError, LookupError):
            continue
        # chardet often calls Hungarian text Latin-1; ő/ű then show up as õ/û. Prefer cp1250.
        if candidate not in {"cp1250", "windows-1250", "iso-8859-2"} and re.search("[õûÕÛ]", text):
            try:
                return content.decode("cp1250"), "cp1250"
            except UnicodeDecodeError:
                pass
        return text, candidate
    return content.decode("latin-1", errors="replace"), "latin-1"


def sniff_delimiter(text: str) -> str:
    lines = [line for line in text.splitlines() if line.strip()][:60]
    if not lines:
        return ","
    best, best_score = ",", -1.0
    for candidate in DELIMITER_CANDIDATES:
        widths = [len(row) for row in csv.reader(lines, delimiter=candidate)]
        counts = Counter(width for width in widths if width > 1)
        if not counts:
            continue
        width, frequency = counts.most_common(1)[0]
        score = frequency / len(widths) + min(width, 30) / 1000
        if score > best_score:
            best, best_score = candidate, score
    return best


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        # repr() gives the shortest string that round-trips the stored IEEE value.
        text = repr(value)
        return text[:-2] if text.endswith(".0") else text
    if isinstance(value, (int, Decimal)):
        return str(value)
    return str(value).strip()


def read_xlsx(content: bytes) -> Table:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell_to_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            if any(_filled(row) for row in rows):
                return Table(rows=rows, kind="xlsx", sheet_name=sheet.title, typed_cells=True)
    finally:
        workbook.close()
    raise UnsupportedFileError("The workbook has no data.")


def read_xls(content: bytes) -> Table:
    book = xlrd.open_workbook(file_contents=content)
    for sheet in book.sheets():
        rows: list[list[str]] = []
        for row_index in range(sheet.nrows):
            cells: list[str] = []
            for cell in sheet.row(row_index):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    cells.append(_cell_to_text(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    cells.append(_cell_to_text(float(cell.value)))
                else:
                    cells.append(_cell_to_text(cell.value))
            rows.append(cells)
        if any(_filled(row) for row in rows):
            return Table(rows=rows, kind="xls", sheet_name=sheet.name, typed_cells=True)
    raise UnsupportedFileError("The workbook has no data.")


class _TableExtractor(HTMLParser):
    """Pulls the largest <table> out of an HTML page (banks love exporting HTML as '.xls')."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self._depth += 1
            self.tables.append([])
        elif tag == "tr" and self._depth:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self.tables:
                self.tables[-1].append(self._row)
            self._row = None
        elif tag == "table" and self._depth:
            self._depth -= 1

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def read_html_table(text: str, encoding: str) -> Table:
    extractor = _TableExtractor()
    extractor.feed(text)
    tables = [table for table in extractor.tables if table]
    if not tables:
        raise UnsupportedFileError("No table found in the HTML export.")
    rows = max(tables, key=len)
    width = max(len(row) for row in rows)
    return Table(rows=[row + [""] * (width - len(row)) for row in rows], kind="html", encoding=encoding)
